"""Crea las categorías, publica las imágenes en Odoo y genera Productos.xlsx.

Las imágenes se suben como adjuntos públicos de Odoo para que la columna
"Imagen" del Excel sea una URL que el importador de Odoo pueda descargar.
"""

import os

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font, PatternFill
from PIL import Image

from datos_productos import CATEGORIA_PADRE, PRODUCTOS, ean13
from odoo_api import adjunto_publico, call

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
IMG_DIR = os.path.join(ROOT, "assets", "productos")
SALIDA = os.path.join(ROOT, "Productos.xlsx")


def categoria(nombre, padre=False):
    ids = call("product.category", "search", domain=[["name", "=", nombre], ["parent_id", "=", padre]])
    return ids[0] if ids else call("product.category", "create", vals_list=[{"name": nombre, "parent_id": padre}])[0]


padre = categoria(CATEGORIA_PADRE)
for sub in sorted({p[3] for p in PRODUCTOS}):
    categoria(sub, padre)

wb = Workbook()
ws = wb.active
ws.title = "Productos"
encabezados = ["Referencia interna", "Nombre", "Tipo de producto", "Rastrear inventario",
               "Categoría del producto", "Precio de venta", "Costo", "Código de barras", "Imagen",
               "Vista previa"]
ws.append(encabezados)
for c in ws[1]:
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor="1F6F50")
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

thumbs = os.path.join(IMG_DIR, "_thumbs")
os.makedirs(thumbs, exist_ok=True)
for i, (ref, nombre, tipo, cat, precio, costo) in enumerate(PRODUCTOS, start=1):
    ruta = os.path.join(IMG_DIR, f"P{i:02d}.jpg")
    ws.append([ref, nombre, tipo, True, f"{CATEGORIA_PADRE} / {cat}", precio, costo,
               ean13(i), adjunto_publico(f"{ref}.jpg", ruta), None])
    fila = i + 1
    ws.row_dimensions[fila].height = 62
    for col in ("F", "G"):
        ws[f"{col}{fila}"].number_format = '"Q"#,##0.00'
    ws[f"H{fila}"].number_format = "@"
    mini = os.path.join(thumbs, f"P{i:02d}.png")
    Image.open(ruta).resize((80, 80)).save(mini)
    ws.add_image(XLImage(mini), f"J{fila}")

for col, ancho in zip("ABCDEFGHIJ", (14, 44, 15, 11, 38, 13, 11, 17, 60, 13)):
    ws.column_dimensions[col].width = ancho
ws.freeze_panes = "A2"
wb.save(SALIDA)
print("Generado:", SALIDA)
